"""
Reads test_data/dmrv_test_data.xlsx and loads all records into the running API.

Usage:
  python scripts/load_test_data.py [--base-url http://localhost:8000]

What it does (in order):
  1. Registers + logs in village chairmen as users
  2. Creates a project (Location) for each village chairman
  3. Registers + logs in all 20 customers as users
  4. Prints a device reference table (devices aren't API objects; they are
     used as metadata when submitting measurements)
  5. Prints a summary of loaded IDs for copy-paste into further tests
"""

import argparse
import json
import sys
import os

import httpx
from openpyxl import load_workbook

XLSX = os.path.join(os.path.dirname(__file__), "..", "test_data", "dmrv_test_data.xlsx")


def sheet_to_dicts(wb, sheet_name):
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h) for h in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)]


def register_user(client, base, row, extra=None):
    payload = {
        "email":       row["email"],
        "full_name":   row["full_name"],
        "password":    row["password"],
        "phone_number": row.get("phone_number"),
        "organization": row.get("organization"),
    }
    r = client.post(f"{base}/api/v1/auth/register", json=payload)
    if r.status_code == 201:
        return r.json()
    if r.status_code == 400 and "already" in r.text.lower():
        # already exists — just login
        pass
    else:
        print(f"  WARN register {row['email']}: {r.status_code} {r.text[:120]}")
    return None


def login_user(client, base, email, password):
    r = client.post(f"{base}/api/v1/auth/login", json={"email": email, "password": password})
    if r.status_code == 200:
        return r.json()["access_token"]
    print(f"  ERROR login {email}: {r.status_code} {r.text[:120]}")
    return None


def create_project(client, base, token, location):
    headers = {"Authorization": f"Bearer {token}"}
    payload = {
        "name":          location["name"],
        "description":   location["description"],
        "project_type":  location["project_type"],
        "country":       location["country"],
        "latitude":      float(location["latitude"] or 0),
        "longitude":     float(location["longitude"] or 0),
        "area_hectares": float(location["area_hectares"] or 0),
        "methodology":   location["methodology"],
        "standard":      location["standard"],
    }
    r = client.post(f"{base}/api/v1/projects/", json=payload, headers=headers)
    if r.status_code == 201:
        return r.json()
    print(f"  WARN create project '{location['name']}': {r.status_code} {r.text[:120]}")
    return None


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://localhost:8000")
    args = parser.parse_args()
    base = args.base_url.rstrip("/")

    if not os.path.exists(XLSX):
        print(f"ERROR: {XLSX} not found.\nRun: python scripts/generate_test_data.py first.")
        sys.exit(1)

    wb = load_workbook(XLSX)
    locations   = sheet_to_dicts(wb, "Locations")
    chairmen    = sheet_to_dicts(wb, "Village_Chairmen")
    customers   = sheet_to_dicts(wb, "Customers")
    devices     = sheet_to_dicts(wb, "Devices")

    loc_map = {loc["location_id"]: loc for loc in locations}
    project_id_map = {}   # location_id → project_id
    token_map = {}        # email → token

    print(f"\n{'='*60}")
    print(f"  dMRV PayGo — Test Data Loader")
    print(f"  Target: {base}")
    print(f"{'='*60}\n")

    client = httpx.Client(timeout=30)

    # ── 1. Village Chairmen ──────────────────────────────────────────────────
    print("[ 1/4 ] Loading Village Chairmen …")
    for vc in chairmen:
        register_user(client, base, vc)
        token = login_user(client, base, vc["email"], vc["password"])
        if not token:
            continue
        token_map[vc["email"]] = token

        loc = loc_map.get(vc["location_id"])
        if loc:
            proj = create_project(client, base, token, loc)
            if proj:
                project_id_map[vc["location_id"]] = proj["id"]
                print(f"  ✓ {vc['full_name']} — project '{loc['name']}' → {proj['id']}")
            else:
                print(f"  ~ {vc['full_name']} — project may already exist")

    # ── 2. Customers ─────────────────────────────────────────────────────────
    print(f"\n[ 2/4 ] Loading Customers …")
    loaded_customers = []
    for cust in customers:
        register_user(client, base, cust)
        token = login_user(client, base, cust["email"], cust["password"])
        if token:
            token_map[cust["email"]] = token
            loaded_customers.append(cust)
            print(f"  ✓ {cust['full_name']} ({cust['role']}) — {cust['email']}")
        else:
            print(f"  ✗ {cust['full_name']} — login failed")

    client.close()

    # ── 3. Device reference table ─────────────────────────────────────────────
    print(f"\n[ 3/4 ] Device Reference Table (use in measurement metadata):")
    print(f"  {'device_id':<10} {'device_name':<28} {'type':<22} {'serial':<16} {'location_id'}")
    print(f"  {'-'*10} {'-'*28} {'-'*22} {'-'*16} {'-'*8}")
    for d in devices:
        print(f"  {d['device_id']:<10} {d['device_name']:<28} {d['device_type']:<22} {d['serial_number']:<16} {d['location_id']}")

    # ── 4. Summary ────────────────────────────────────────────────────────────
    print(f"\n[ 4/4 ] Summary:")
    print(f"  Village Chairmen loaded : {len([c for c in chairmen if c['email'] in token_map])}/{len(chairmen)}")
    print(f"  Customers loaded        : {len(loaded_customers)}/{len(customers)}")
    print(f"  Projects created        : {len(project_id_map)}")
    print(f"  Devices in reference    : {len(devices)}")

    print(f"\n  Project ID Map (location_id → API project_id):")
    for loc_id, proj_id in project_id_map.items():
        print(f"    {loc_id} → {proj_id}")

    print(f"\n  Sample measurement payload (copy & paste to POST /api/v1/projects/<project_id>/measurements/):")
    sample_loc = "LOC-001"
    sample_proj = project_id_map.get(sample_loc, "<project_id>")
    print(json.dumps({
        "measurement_type": "carbon_sequestration",
        "value": 12.5,
        "unit": "tCO2e",
        "measurement_date": "2024-06-01T08:00:00Z",
        "notes": "Automated reading from KaruraCO2Sensor-01",
        "metadata": {
            "device_id": "DEV-002",
            "device_name": "KaruraCO2Sensor-01",
            "serial_number": "CO2-KR-0002",
            "location_id": "LOC-001",
        },
    }, indent=4))

    print(f"\n{'='*60}")
    print("  Done. Open http://localhost:8000/docs to explore the API.")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    main()
